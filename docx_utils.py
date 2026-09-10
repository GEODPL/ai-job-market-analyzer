
import io
from docx import Document
from docx.shared import Pt, Inches, RGBColor


def create_docx(title: str, content: str, doc_type: str = "Document") -> io.BytesIO:
    doc = Document()
    for s in doc.sections:
        s.top_margin = Inches(1)
        s.bottom_margin = Inches(1)
        s.left_margin = Inches(1)
        s.right_margin = Inches(1)

    p_title = doc.add_paragraph()
    r_title = p_title.add_run(title)
    r_title.font.name = 'Calibri'
    r_title.font.size = Pt(17)
    r_title.font.bold = True
    r_title.font.color.rgb = RGBColor(30, 58, 138)
    p_title.paragraph_format.space_after = Pt(2)

    p_sub = doc.add_paragraph()
    r_sub = p_sub.add_run(f"AI Generated {doc_type} | Greek AI Job Market Analyzer")
    r_sub.font.name = 'Calibri'
    r_sub.font.size = Pt(9.5)
    r_sub.font.color.rgb = RGBColor(100, 116, 139)
    p_sub.paragraph_format.space_after = Pt(16)

    for line in content.split('\n'):
        line_clean = line.strip()
        if not line_clean:
            continue
        p = doc.add_paragraph()
        p.paragraph_format.line_spacing = 1.15
        p.paragraph_format.space_after = Pt(5)

        if line_clean.startswith(('-', '•', '*')) or (
            line_clean[0].isdigit() and line_clean[1:3] in ('. ', ') ')
        ):
            p.paragraph_format.left_indent = Inches(0.25)

        run = p.add_run(line_clean)
        run.font.name = 'Calibri'
        run.font.size = Pt(11)
        run.font.color.rgb = RGBColor(30, 41, 59)

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf